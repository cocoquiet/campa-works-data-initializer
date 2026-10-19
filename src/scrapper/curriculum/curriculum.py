import httpx
import json
import openpyxl

from scrapper.master_course.master_course import (
    excel_file_path,
    fetch_master_courses,
)

curriculum_url = (
    "https://knuin.knu.ac.kr/public/web/stddm/curse/educourse/eduCour/searchEduCourList"
)

years = range(2020, 2027)
SEMESTER_CODES = {
    "FIRST": ("STCU000600001", "STCU000600003"),
    "SECOND": ("STCU000600002",),
}

request_payload = {
    "isApi": "Y",
    "search": {
        "dprtnCd": "",
        "dprtnNm": "",
        "forMdlMgmntYn": "Y",
        "frstNcrYear": 2026,
        "isApi": "Y",
        "mdlSctinCd": "STCU001800001",
        "trgtYrsf": 0,
    },
}


def fetch_excel_curriculums():
    major_codes = {
        "인공지능시스템학과": "A701",
        "컴퓨터학부(심화컴퓨팅전공)": "A70201",
        "컴퓨터학부(첨단AI컴퓨팅전공)": "A70202",
    }

    workbook = openpyxl.load_workbook(excel_file_path, data_only=True)
    records = []

    try:
        for sheet_name, major_code in major_codes.items():
            sheet = workbook[sheet_name]
            grade = None
            reading_model = False

            for row in sheet.iter_rows(max_col=9, values_only=True):
                values = row[1:9]
                grade_text = str(values[0]).strip()

                if "표준개설" in grade_text:
                    reading_model = True
                    continue
                if not reading_model:
                    continue
                if "전공이수학점" in str(values):
                    break

                if grade_text == "학년":
                    grade = None
                elif grade_text in ["1학년", "2학년", "3학년", "4학년"]:
                    grade = int(grade_text[0])

                if grade is None:
                    continue

                if values[1] and values[2]:
                    records.append(
                        {
                            "major_code": major_code,
                            "year": 2027,
                            "semester_": "FIRST",
                            "grade": grade,
                        }
                    )

                if values[4] and values[5]:
                    records.append(
                        {
                            "major_code": major_code,
                            "year": 2027,
                            "semester_": "SECOND",
                            "grade": grade,
                        }
                    )
    finally:
        workbook.close()
    return records


def post_single_curriculum(data):
    response = httpx.post(
        "http://localhost:8080/api/curriculums",
        json=data,
        timeout=20.0,
    )

    if response.status_code != 201:
        print(
            f"Failed to post curriculum data for {data}. "
            f"Status code: {response.status_code}"
        )
    else:
        print(f"Successfully posted data for {data}.")

    return response.status_code in [201, 409]


def get_curriculum_grades(major_code, records, year, semester_codes):
    if str(major_code).startswith("2"):
        return {1}

    grades = set()
    for record in records:
        if (
            str(record.get("dprtnCd")) != str(major_code)
            or int(record.get("trgtYrsf")) != year
            or record.get("estblSmstrCd") not in semester_codes
        ):
            continue

        try:
            grade = int(record.get("estblGrade"))
        except TypeError, ValueError:
            continue

        if grade in range(1, 5):
            grades.add(grade)

    return grades


def post_excel_curriculums(majors, semesters):
    records = fetch_excel_curriculums()
    posted = set()

    for record in records:
        for major in majors:
            if major["major_code"] != record["major_code"]:
                continue

            for semester in semesters:
                if (
                    semester["year"] != record["year"]
                    or semester["semester_"] != record["semester_"]
                ):
                    continue

                data = {
                    "major_id": major["id"],
                    "semester_id": semester["id"],
                    "grade": record["grade"],
                }

                key = (
                    data["major_id"],
                    data["semester_id"],
                    data["grade"],
                )

                if key in posted:
                    continue
                if post_single_curriculum(data):
                    posted.add(key)


def post_curriculums():

    majors = httpx.get("http://localhost:8080/api/majors").json()
    semesters = httpx.get("http://localhost:8080/api/semesters").json()

    for major in majors:
        for semester in semesters:
            if semester["year"] not in years:
                continue

            semester_codes = SEMESTER_CODES.get(semester["semester_"])

            if semester_codes is None:
                continue

            request_payload["search"]["dprtnCd"] = major["major_code"]
            request_payload["search"]["trgtYrsf"] = semester["year"]

            response = httpx.post(
                curriculum_url,
                json=request_payload,
                timeout=20.0,
            )
            response.raise_for_status()
            records = response.json()["data"]

            grades = get_curriculum_grades(
                major["major_code"],
                records,
                semester["year"],
                semester_codes,
            )

            for grade in sorted(grades):
                post_single_curriculum(
                    {
                        "major_id": major["id"],
                        "semester_id": semester["id"],
                        "grade": grade,
                    }
                )
    post_excel_curriculums(majors, semesters)


# wb = openpyxl.Workbook()
# for major_code, major_name in majors.items():
#     print(f"\nProcessing {major_name}({major_code})...")
#     ws = wb.create_sheet(title=major_name)
#     ws.append(
#         [
#             "년도",
#             "학년",
#             "학기",
#             "교과목코드",
#             "교과목명",
#             "학점",
#             "이론",
#             "실습",
#             "전공필수구분",
#             "표준표기",
#         ]
#     )

#     for year in years:
#         print(f"  {year}...", end="")
#         while True:
#             try:
#                 request_payload["search"]["dprtnCd"] = major_code
#                 request_payload["search"]["trgtYrsf"] = year

#                 response = httpx.post(
#                     curriculum_url, json=request_payload, timeout=20.0
#                 )
#                 data = json.loads(response.text)["data"]
#                 print(f"Retrieved {len(data)} records.")
#                 break
#             except httpx.TimeoutException as e:
#                 continue
#             except Exception as e:
#                 print(f"An error occurred for year {year}: {e}")
#                 break

#         try:
#             for item in data:
#                 ws.append(
#                     [
#                         item["trgtYrsf"],
#                         item["estblGrade"],
#                         item["estblSmstrCd"][-1],
#                         item["sbjetCd"],
#                         item["sbjetNm"],
#                         item["crdit"],
#                         item["thryTime"],
#                         item["prctsTime"],
#                         "전공필수" if item["mnmjMjrYn"] == "Y" else "",
#                         f"{item['sbjetNm']}-{major_name}-{year}-{item['estblGrade']}-{item['estblSmstrCd'][-1]}-{item['crditSystem']}",
#                     ]
#                 )
#             wb.save("curriculum.xlsx")
#         except Exception as e:
#             print(f"\tAn error occurred while saving data for year {year}: {e}")
#         else:
#             print(f"\tSaved data for year {year} to curriculum.xlsx.")

if __name__ == "__main__":
    post_curriculums()
